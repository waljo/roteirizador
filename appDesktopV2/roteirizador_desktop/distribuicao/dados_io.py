from __future__ import annotations

from datetime import time

import openpyxl

from .models import DadosRow, FilledRow

_C_TP = 1
_C_NOTA = 2
_C_ITEM = 3
_C_SUBITEM = 4
_C_DESC = 5
_C_NAME = 6
_C_DATE = 7
_C_ORIGEM = 8
_C_DESTINO = 9
_C_EMBARCACAO = 10
_C_HORARIO = 11
_C_N_VIAGEM = 12
_C_TIPO = 13


def _dados_sheet(wb):
    """A aba `Dados`, e nao a que estava selecionada quando o arquivo foi salvo.

    A pasta de trabalho do operador tem tres abas (`TD`, `Dados`, `Planilha1`) e o `wb.active`
    e so a que ficou em foco no ultimo salvamento. Achado no arquivo de 21/08, salvo com a
    `TD` (uma tabela dinamica) em foco: o `read_dados` devolvia **zero linhas em silencio** —
    o operador processa e nao aparece nada — e o `write_dados` gravaria as quatro colunas
    dentro da dinamica, nas linhas de outra coisa.

    As duas funcoes tem de usar a mesma aba, senao a gravacao vai para outro lugar que a
    leitura.
    """
    for nome in wb.sheetnames:
        if nome.strip().casefold() == "dados":
            return wb[nome]
    return wb.active


def read_dados(path: str) -> list[DadosRow]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = _dados_sheet(wb)
    rows: list[DadosRow] = []

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def get(col: int):
            return row[col - 1] if len(row) >= col else None

        name = get(_C_NAME)
        if not name:
            continue

        horario = get(_C_HORARIO)
        if not isinstance(horario, time):
            horario = None

        nota_val = get(_C_NOTA)
        item_val = get(_C_ITEM)
        n_viagem_val = get(_C_N_VIAGEM)

        rows.append(DadosRow(
            excel_row=excel_row,
            tp=get(_C_TP),
            nota=str(nota_val) if nota_val is not None else None,
            item=str(item_val) if item_val is not None else None,
            subitem=str(get(_C_SUBITEM)) if get(_C_SUBITEM) is not None else None,
            desc=get(_C_DESC),
            name=str(name),
            date_str=str(get(_C_DATE)) if get(_C_DATE) is not None else None,
            origem_raw=str(get(_C_ORIGEM)) if get(_C_ORIGEM) is not None else None,
            destino_raw=str(get(_C_DESTINO)) if get(_C_DESTINO) is not None else None,
            embarcacao=str(get(_C_EMBARCACAO)) if get(_C_EMBARCACAO) is not None else None,
            horario=horario,
            n_viagem=int(n_viagem_val) if n_viagem_val is not None else None,
            tipo_viagem=str(get(_C_TIPO)).strip() if get(_C_TIPO) is not None else None,
        ))

    wb.close()
    return rows


def write_dados(path: str, filled_rows: list[FilledRow]) -> None:
    wb = openpyxl.load_workbook(path)
    ws = _dados_sheet(wb)

    for fr in filled_rows:
        if fr.status == "already_filled":
            continue
        r = fr.dados_row.excel_row

        # Sem embarcacao a linha nao esta programada, e ai as quatro colunas saem em branco —
        # inclusive o TIPO. Gravar so o tipo deixava a linha parecendo meio programada e o
        # colega que gera os manifestos vinha perguntar se faltava programar aqueles pax.
        # A regra olha a embarcacao, e nao o status, para nunca descartar uma embarcacao que
        # o operador digitou a mao na tabela.
        if not fr.embarcacao:
            for col in (_C_EMBARCACAO, _C_HORARIO, _C_N_VIAGEM, _C_TIPO):
                ws.cell(row=r, column=col).value = None
            continue

        ws.cell(row=r, column=_C_EMBARCACAO).value = fr.embarcacao
        ws.cell(row=r, column=_C_HORARIO).value = fr.horario
        ws.cell(row=r, column=_C_N_VIAGEM).value = fr.n_viagem
        ws.cell(row=r, column=_C_TIPO).value = (fr.tipo_viagem + " ") if fr.tipo_viagem else None

    wb.save(path)
    wb.close()
